from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path
from subprocess import CompletedProcess
from unittest.mock import patch

from docx import Document
from PIL import Image

from writing_brain.config import AppConfig
from writing_brain.image_supply import GENERATED_INDEX_FILENAME, PUBLIC_INDEX_FILENAME
from writing_brain.publish import (
    _call_gemini_native_image_api,
    _render_image_retry_actions,
    _resolve_image_runtime,
    build_publish_pack,
    render_packy_images,
)


MINI_PNG = (
    b"\x89PNG\r\n\x1a\n"
    b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
    b"\x00\x00\x00\x0cIDAT\x08\x99c\xf8\xcf\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


def _make_image_bytes(size: tuple[int, int], *, image_format: str = "PNG", color: tuple[int, int, int] = (80, 120, 180)) -> bytes:
    image = Image.new("RGB", size, color)
    output = io.BytesIO()
    image.save(output, format=image_format)
    return output.getvalue()


class PublishTests(unittest.TestCase):
    def test_build_publish_pack_creates_expected_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "image-plans").mkdir()
            (root / "image-plans" / "img.md").write_text(
                "---\n"
                "id: imageplan-101\n"
                "name: 观点型-概念图加结构图\n"
                "---\n\n"
                "## 封面图\n\n"
                "- 风格：媒体感概念图，克制。\n",
                encoding="utf-8",
            )
            config = AppConfig(data_dir=root)
            output_dir = root / "out" / "公众号"
            generated_dir = output_dir / "图片" / "已生成"
            public_dir = output_dir / "图片" / "公开来源"
            generated_dir.mkdir(parents=True)
            public_dir.mkdir(parents=True)
            for name in ["封面", "配图-02"]:
                (generated_dir / f"{name}.png").write_bytes(MINI_PNG)
            for name in ["配图-01", "配图-03"]:
                (public_dir / f"{name}.png").write_bytes(MINI_PNG)
            (output_dir / "图片" / GENERATED_INDEX_FILENAME).write_text(
                json.dumps(
                    {
                        "封面": {"path": str(generated_dir / "封面.png"), "provider": "packyapi", "model": "nano2", "publishable": True},
                        "配图-02": {"path": str(generated_dir / "配图-02.png"), "provider": "packyapi", "model": "nano2", "publishable": True},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (output_dir / "图片" / PUBLIC_INDEX_FILENAME).write_text(
                json.dumps(
                    {
                        "配图-01": {
                            "path": str(public_dir / "配图-01.png"),
                            "source_url": "https://commons.wikimedia.org/wiki/File:one",
                            "license": "CC BY-SA 4.0",
                            "author": "Author One",
                            "publishable": True,
                        },
                        "配图-03": {
                            "path": str(public_dir / "配图-03.png"),
                            "source_url": "https://commons.wikimedia.org/wiki/File:three",
                            "license": "CC BY-SA 4.0",
                            "author": "Author Three",
                            "publishable": True,
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = build_publish_pack(
                {
                    "title": "中国AI行业最危险的，不是落后，而是开始适应落后",
                    "topic": "中国AI行业最危险的，不是落后，而是开始适应落后",
                    "platform": "wechat",
                    "output_dir": str(output_dir),
                    "article_markdown": (
                        "# 中国AI行业最危险的，不是落后，而是开始适应落后\n\n"
                        "中国 AI 行业现在最危险的，不是模型暂时落后，而是开始适应落后。\n\n"
                        "版本升级不是名字变化，而是能力边界变化。\n\n"
                        "长期不用最强模型的人，会系统性误判未来。\n\n"
                        "入口重要，但模型上限决定谁定义明天。"
                    ),
                },
                config,
            )

            self.assertEqual(result["image_task_count"], 2)
            self.assertEqual(result["images_present"], 4)
            self.assertEqual(result["image_review_report"]["decision"], "pass")
            task_path = output_dir / "图片" / "生成任务.json"
            placements_path = output_dir / "图片" / "插图位置.json"
            leads_path = output_dir / "图片" / "公开图片线索.md"
            pure_docx = output_dir / "公众号版-可直接发布-纯文本可复制.docx"
            rich_docx = output_dir / "公众号版-图文可发布.docx"
            self.assertTrue(task_path.exists())
            self.assertTrue(placements_path.exists())
            self.assertTrue(leads_path.exists())
            self.assertTrue(pure_docx.exists())
            self.assertTrue(rich_docx.exists())

            tasks = json.loads(task_path.read_text(encoding="utf-8"))
            self.assertEqual(tasks[0]["filename"], "封面")
            self.assertEqual(tasks[1]["filename"], "配图-02")

            rich = Document(rich_docx)
            self.assertEqual(len(rich.inline_shapes), 4)
            rich_text = "\n".join(paragraph.text for paragraph in rich.paragraphs)
            self.assertNotIn("封面图：", rich_text)
            self.assertNotIn("配图-02", rich_text)

    def test_build_publish_pack_marks_image_review_revise_when_assets_missing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "image-plans").mkdir()
            config = AppConfig(data_dir=root)
            output_dir = root / "out" / "公众号"

            result = build_publish_pack(
                {
                    "title": "中国AI行业最危险的，不是落后，而是开始适应落后",
                    "topic": "中国AI行业最危险的，不是落后，而是开始适应落后",
                    "platform": "wechat",
                    "output_dir": str(output_dir),
                    "article_markdown": (
                        "# 中国AI行业最危险的，不是落后，而是开始适应落后\n\n"
                        "中国 AI 行业现在最危险的，不是模型暂时落后，而是开始适应落后。\n\n"
                        "版本升级不是名字变化，而是能力边界变化。\n\n"
                        "长期不用最强模型的人，会系统性误判未来。\n\n"
                        "入口重要，但模型上限决定谁定义明天。"
                    ),
                },
                config,
            )

            self.assertEqual(result["image_review_report"]["decision"], "revise")
            self.assertTrue(any("图片素材未补齐" in item["summary"] for item in result["image_review_report"]["top_issues"]))
            self.assertTrue(any("补齐缺失图片" in action for action in result["image_review_report"]["required_actions"]))

    def test_build_publish_pack_renders_docx_formatting_from_markdown_structure(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "image-plans").mkdir()
            config = AppConfig(data_dir=root)
            output_dir = root / "out" / "公众号"

            build_publish_pack(
                {
                    "title": "AI 工作流不是拼提示词",
                    "topic": "AI 工作流不是拼提示词",
                    "platform": "wechat",
                    "output_dir": str(output_dir),
                    "article_markdown": (
                        "# AI 工作流不是拼提示词\n\n"
                        "先说结论，真正重要的是把流程拆清楚。\n\n"
                        "## 为什么会失控\n\n"
                        "很多团队把 **不确定性** 直接送进高成本环节，所以返工会越来越多。\n\n"
                        "- 先锁角色\n"
                        "- 再锁场景\n\n"
                        "### 落地动作\n\n"
                        "1. 先做低成本试错\n"
                        "2. 再进入正式生成"
                    ),
                },
                config,
            )

            pure_docx = output_dir / "公众号版-可直接发布-纯文本可复制.docx"
            doc = Document(pure_docx)
            paragraphs = [(paragraph.text, paragraph.style.name) for paragraph in doc.paragraphs if paragraph.text.strip()]

            self.assertIn(("为什么会失控", "Heading 1"), paragraphs)
            self.assertIn(("落地动作", "Heading 2"), paragraphs)
            self.assertIn(("先锁角色", "List Bullet"), paragraphs)
            self.assertIn(("先做低成本试错", "List Number"), paragraphs)

            target = next(paragraph for paragraph in doc.paragraphs if "不确定性" in paragraph.text)
            self.assertTrue(any(run.bold and "不确定性" in run.text for run in target.runs))

    def test_build_publish_pack_normalizes_existing_assets_and_removes_duplicate_variants(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "image-plans").mkdir()
            config = AppConfig(data_dir=root)
            output_dir = root / "out" / "公众号"
            generated_dir = output_dir / "图片" / "已生成"
            public_dir = output_dir / "图片" / "公开来源"
            generated_dir.mkdir(parents=True)
            public_dir.mkdir(parents=True)

            (generated_dir / "封面.png").write_bytes(_make_image_bytes((320, 320)))
            (generated_dir / "配图-02.png").write_bytes(_make_image_bytes((600, 600)))
            (public_dir / "配图-01.png").write_bytes(_make_image_bytes((500, 300)))
            (public_dir / "配图-01.jpg").write_bytes(_make_image_bytes((300, 500), image_format="JPEG"))
            (public_dir / "配图-03.png").write_bytes(_make_image_bytes((420, 420)))

            build_publish_pack(
                {
                    "title": "AI 工作流不是拼提示词",
                    "topic": "AI 工作流不是拼提示词",
                    "platform": "wechat",
                    "output_dir": str(output_dir),
                    "article_markdown": (
                        "# AI 工作流不是拼提示词\n\n"
                        "先说结论，真正重要的是把流程拆清楚。\n\n"
                        "很多团队把不确定性直接送进高成本环节，所以返工会越来越多。\n\n"
                        "先锁角色，再锁场景，再进入正式生成。"
                    ),
                },
                config,
            )

            self.assertFalse((public_dir / "配图-01.jpg").exists())
            with Image.open(generated_dir / "封面.png") as image:
                self.assertEqual(image.size, (1600, 900))
            with Image.open(generated_dir / "配图-02.png") as image:
                self.assertEqual(image.size, (1600, 1200))
            with Image.open(public_dir / "配图-01.png") as image:
                self.assertEqual(image.size, (1600, 1200))
            with Image.open(public_dir / "配图-03.png") as image:
                self.assertEqual(image.size, (1600, 1200))

    def test_render_packy_images_waits_for_api_key(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tasks_path = Path(tmp) / "生成任务.json"
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.dict("os.environ", {}, clear=True):
                result = render_packy_images({"tasks_path": str(tasks_path)})
            self.assertEqual(result["status"], "awaiting_api_key")

    def test_render_packy_images_writes_generated_index(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.dict("os.environ", {"PACKYAPI_IMAGE_API_KEY": "test-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                return_value={"status": "ok", "bytes": MINI_PNG, "extension": "png", "mime_type": "image/png"},
            ) as mock_call:
                result = render_packy_images({"tasks_path": str(tasks_path)})

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["resolved_provider"], "packyapi")
            self.assertEqual(result["resolved_model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(result["resolved_base_url"], "https://www.packyapi.com")
            self.assertEqual(mock_call.call_args.kwargs["base_url"], "https://www.packyapi.com")
            index_path = image_dir / GENERATED_INDEX_FILENAME
            self.assertTrue(index_path.exists())
            index = json.loads(index_path.read_text(encoding="utf-8"))
            self.assertEqual(index["封面"]["model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(index["封面"]["provider"], "packyapi")
            self.assertTrue(index["封面"]["publishable"])

    def test_render_packy_images_does_not_reuse_generic_packy_chat_key(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            tasks_path = Path(tmp) / "生成任务.json"
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.dict(
                "os.environ",
                {"PACKYAPI_API_KEY": "chat-key", "PACKYAPI_BASE_URL": "https://www.packyapi.com/v1"},
                clear=True,
            ):
                result = render_packy_images({"tasks_path": str(tasks_path)})

            self.assertEqual(result["status"], "awaiting_api_key")
            self.assertEqual(result["resolved_provider"], "ikun")
            self.assertTrue(any("IKUN_IMAGE_API_KEY" in item for item in result["recommended_next_actions"]))

    def test_gemini_native_image_api_strips_v1_before_v1beta(self) -> None:
        with patch(
            "writing_brain.publish.subprocess.run",
            return_value=CompletedProcess(args=[], returncode=0, stdout='{"candidates":[]}', stderr=""),
        ) as mock_run:
            _call_gemini_native_image_api(
                prompt="测试图片",
                model="gemini-3.1-flash-image-preview",
                api_key="test-key",
                aspect_ratio="4:3",
                image_size="1K",
                base_url="https://api.ikuncode.cc/v1",
            )

        command = mock_run.call_args.args[0]
        self.assertEqual(command[4], "https://api.ikuncode.cc/v1beta/models/gemini-3.1-flash-image-preview:generateContent")

    def test_gemini_native_image_api_reports_empty_candidate_parts(self) -> None:
        with patch(
            "writing_brain.publish.subprocess.run",
            return_value=CompletedProcess(
                args=[],
                returncode=0,
                stdout='{"candidates":[{"finishReason":"STOP","content":{"parts":[]}}]}',
                stderr="",
            ),
        ):
            result = _call_gemini_native_image_api(
                prompt="测试图片",
                model="gemini-3.1-flash-image-preview",
                api_key="test-key",
                aspect_ratio="4:3",
                image_size="1K",
                base_url="https://api.ikuncode.cc/v1",
            )

        self.assertEqual(result["status"], "error")
        self.assertIn("empty_candidate_parts", result["error"])
        self.assertIn("finish_reason=STOP", result["error"])

    def test_explicit_gemini_preview_model_with_ikun_key_stays_on_ikun(self) -> None:
        with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "test-key"}, clear=True):
            runtime = _resolve_image_runtime({"image_model": "gemini-3-pro-image-preview"})

        self.assertEqual(runtime.provider, "ikun")
        self.assertEqual(runtime.base_url, "https://api.ikuncode.cc/v1")
        self.assertEqual(runtime.model, "gemini-3-pro-image-preview")

    def test_render_packy_images_retries_same_model_before_success(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            call_models: list[str] = []

            def _mock_call(*, model: str, **kwargs):
                call_models.append(model)
                if len(call_models) < 3:
                    return {"status": "error", "error": '{"code":"model_not_found","message":"No available channel"}'}
                return {"status": "ok", "bytes": MINI_PNG, "extension": "png", "mime_type": "image/png"}

            with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "test-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                side_effect=_mock_call,
            ):
                result = render_packy_images({"tasks_path": str(tasks_path)})

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["configured_model"], "nano2")
            self.assertEqual(result["resolved_model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(call_models, ["gemini-3.1-flash-image-preview"] * 3)
            index = json.loads((image_dir / GENERATED_INDEX_FILENAME).read_text(encoding="utf-8"))
            self.assertEqual(index["封面"]["model"], "gemini-3.1-flash-image-preview")

    def test_render_packy_images_defaults_to_ikun_nano2(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "test-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                return_value={"status": "ok", "bytes": MINI_PNG, "extension": "png", "mime_type": "image/png"},
            ) as mock_call:
                result = render_packy_images({"tasks_path": str(tasks_path)})

            self.assertEqual(result["status"], "completed")
            self.assertEqual(result["configured_model"], "nano2")
            self.assertEqual(result["resolved_model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(mock_call.call_args.kwargs["model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(mock_call.call_args.kwargs["base_url"], "https://api.ikuncode.cc/v1")
            index = json.loads((image_dir / GENERATED_INDEX_FILENAME).read_text(encoding="utf-8"))
            self.assertEqual(index["封面"]["provider"], "ikun")
            self.assertEqual(index["封面"]["model"], "gemini-3.1-flash-image-preview")

    def test_render_packy_images_normalizes_output_size_and_removes_old_variants(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            output_dir = image_dir / "已生成"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            output_dir.mkdir(parents=True)
            (output_dir / "封面.png").write_bytes(_make_image_bytes((200, 200)))
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "test-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                return_value={
                    "status": "ok",
                    "bytes": _make_image_bytes((400, 300), image_format="JPEG"),
                    "extension": "jpg",
                    "mime_type": "image/jpeg",
                },
            ):
                result = render_packy_images({"tasks_path": str(tasks_path), "output_dir": str(output_dir)})

            self.assertEqual(result["status"], "completed")
            self.assertFalse((output_dir / "封面.png").exists())
            saved = output_dir / "封面.jpg"
            self.assertTrue(saved.exists())
            with Image.open(saved) as image:
                self.assertEqual(image.size, (1600, 900))
            index = json.loads((image_dir / GENERATED_INDEX_FILENAME).read_text(encoding="utf-8"))
            self.assertEqual(index["封面"]["width"], 1600)
            self.assertEqual(index["封面"]["height"], 900)

    def test_render_packy_images_can_fallback_to_backup_channel(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            calls: list[dict[str, str]] = []

            def _mock_call(*, model: str, api_key: str, base_url: str, **kwargs):
                calls.append({"model": model, "api_key": api_key, "base_url": base_url})
                if api_key == "primary-key":
                    return {"status": "error", "error": "empty_candidate_parts: finish_reason=STOP,parts=0", "attempts_used": 3}
                return {"status": "ok", "bytes": MINI_PNG, "extension": "png", "mime_type": "image/png"}

            with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "primary-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                side_effect=_mock_call,
            ):
                result = render_packy_images(
                    {
                        "tasks_path": str(tasks_path),
                        "image_fallback_provider": "packyapi",
                        "image_fallback_api_key": "fallback-key",
                        "image_fallback_base_url": "https://www.packyapi.com",
                    }
                )

            self.assertEqual(result["status"], "completed")
            self.assertTrue(result["fallback_used"])
            self.assertEqual(result["fallback_provider"], "packyapi")
            self.assertEqual(result["results"][0]["provider"], "packyapi")
            self.assertEqual(result["results"][0]["model"], "gemini-3.1-flash-image-preview")
            self.assertEqual(calls[0]["base_url"], "https://api.ikuncode.cc/v1")
            self.assertEqual(calls[-1]["base_url"], "https://www.packyapi.com")
            index = json.loads((image_dir / GENERATED_INDEX_FILENAME).read_text(encoding="utf-8"))
            self.assertEqual(index["封面"]["provider"], "packyapi")
            self.assertEqual(index["封面"]["model"], "gemini-3.1-flash-image-preview")

    def test_render_packy_images_skips_fallback_without_backup_key(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_dir = root / "图片"
            tasks_path = image_dir / "生成任务.json"
            image_dir.mkdir(parents=True)
            tasks_path.write_text(
                json.dumps(
                    [
                        {
                            "filename": "封面",
                            "prompt": "测试图片",
                            "aspect_ratio": "16:9",
                            "image_size": "2K",
                        }
                    ],
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with patch.dict("os.environ", {"IKUN_IMAGE_API_KEY": "primary-key"}, clear=True), patch(
                "writing_brain.publish._call_image_api",
                return_value={"status": "error", "error": "model_not_found", "attempts_used": 3},
            ) as mock_call:
                result = render_packy_images(
                    {
                        "tasks_path": str(tasks_path),
                        "image_fallback_provider": "packyapi",
                        "image_fallback_base_url": "https://www.packyapi.com",
                    }
                )

            self.assertEqual(result["status"], "partial")
            self.assertFalse(result["fallback_used"])
            self.assertEqual(mock_call.call_count, 3)

    def test_render_retry_actions_explain_empty_candidate_parts(self) -> None:
        runtime = _resolve_image_runtime({"image_provider": "ikun"})
        actions = _render_image_retry_actions(
            [{"status": "error", "error": "empty_candidate_parts: finish_reason=STOP,parts=0"}],
            runtime,
            retry_attempts=3,
        )

        self.assertTrue(any("candidates.content.parts 为空" in item for item in actions))


if __name__ == "__main__":
    unittest.main()
