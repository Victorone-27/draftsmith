from __future__ import annotations

import base64
import io
import json
import os
import re
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageOps, UnidentifiedImageError

from .config import AppConfig
from .image_review import build_image_review_report
from .image_supply import GENERATED_INDEX_FILENAME, read_image_index, write_image_index
from .llm import normalize_base_url
from .platforms import PLATFORM_FILE_STEMS, PLATFORM_NAMES, normalize_platform, platform_display_name, platform_file_stem
from .retrieval import load_markdown_items, rank_items
from .text import compact_whitespace, extract_terms, slug_for_filename, split_paragraphs

DEFAULT_IMAGE_MODEL = "nano2"
DEFAULT_PACKY_IMAGE_MODEL = "gemini-3.1-flash-image-preview"
DEFAULT_IMAGE_COUNT = 3
DEFAULT_IKUN_IMAGE_BASE_URL = "https://api.ikuncode.cc"
DEFAULT_PACKY_IMAGE_BASE_URL = "https://www.packyapi.com"
DEFAULT_IMAGE_RETRY_ATTEMPTS = 3
DEFAULT_STANDARD_IMAGE_LONG_EDGE = 1600
IMAGE_MODEL_ALIASES = {
    "nano2": "gemini-3.1-flash-image-preview",
}
STANDARD_IMAGE_TARGETS = {
    "16:9": (1600, 900),
    "4:3": (1600, 1200),
    "3:4": (1200, 1600),
    "1:1": (1600, 1600),
}


@dataclass(frozen=True)
class ImageSlot:
    filename: str
    caption: str
    prompt: str
    aspect_ratio: str
    image_size: str
    source_type: str
    role: str = "supporting_visual"
    allowed_source_types: tuple[str, ...] = ("generated",)
    after_contains: str | None = None


@dataclass(frozen=True)
class ImageRuntime:
    provider: str
    api_key: str
    model: str
    base_url: str
    key_source: str
    model_source: str
    base_url_source: str


def build_publish_pack(payload: dict[str, Any], config: AppConfig) -> dict[str, Any]:
    article_markdown = str(
        payload.get("article_markdown")
        or payload.get("final_article_markdown")
        or payload.get("draft_text")
        or ""
    ).strip()
    if not article_markdown:
        raise ValueError("article_markdown 不能为空")

    context_pack = dict(payload.get("context_pack") or {})
    platform = normalize_platform(payload.get("platform") or context_pack.get("platform") or "wechat")
    topic = str(payload.get("topic") or context_pack.get("topic") or "").strip()
    title = str(payload.get("title") or _extract_title(article_markdown) or topic or "未命名文章").strip()
    image_count = max(3, min(int(payload.get("image_count") or DEFAULT_IMAGE_COUNT), 5))
    output_dir = _resolve_output_dir(payload, config, platform=platform, title=title)
    image_dir = output_dir / "图片"
    generated_dir = image_dir / "已生成"
    public_dir = image_dir / "公开来源"
    generated_dir.mkdir(parents=True, exist_ok=True)
    public_dir.mkdir(parents=True, exist_ok=True)

    image_plan = _select_image_plan(config, payload, context_pack, article_markdown=article_markdown)
    image_brief = dict(payload.get("image_brief") or {})
    if image_brief.get("slots"):
        slots = _build_image_slots_from_brief(
            image_brief=image_brief,
            title=title,
            topic=topic or title,
            article_markdown=article_markdown,
            platform=platform,
        )
    else:
        slots = _build_image_slots(
            title=title,
            topic=topic or title,
            article_markdown=article_markdown,
            platform=platform,
            image_count=image_count,
            image_plan_name=image_plan.title if image_plan else "",
            image_plan_body=image_plan.body if image_plan else "",
        )

    tasks = [
        {
            "filename": slot.filename,
            "source_type": slot.source_type,
            "role": slot.role,
            "allowed_source_types": list(slot.allowed_source_types),
            "aspect_ratio": slot.aspect_ratio,
            "image_size": slot.image_size,
            "prompt": slot.prompt,
        }
        for slot in slots
        if "generated" in slot.allowed_source_types
    ]
    placements = _build_placements(slots)
    source_markdown = _build_source_markdown(
        slots=slots,
        payload=payload,
        image_model=_resolve_image_model(payload),
    )
    public_leads_markdown = _build_public_image_leads(slots)

    tasks_path = image_dir / "生成任务.json"
    placements_path = image_dir / "插图位置.json"
    sources_path = image_dir / "图片与数据来源.md"
    public_leads_path = image_dir / "公开图片线索.md"
    tasks_path.write_text(json.dumps(tasks, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    placements_path.write_text(json.dumps(placements, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    sources_path.write_text(source_markdown + ("\n" if not source_markdown.endswith("\n") else ""), encoding="utf-8")
    public_leads_path.write_text(
        public_leads_markdown + ("\n" if not public_leads_markdown.endswith("\n") else ""),
        encoding="utf-8",
    )
    _normalize_slot_assets(slots, {"generated": generated_dir, "public": public_dir})

    pure_docx_path = output_dir / f"{_platform_file_stem(platform)}-可直接发布-纯文本可复制.docx"
    rich_docx_path = output_dir / f"{_platform_file_stem(platform)}-图文可发布.docx"
    _write_docx(article_markdown, pure_docx_path, placements=None, image_dirs={"generated": generated_dir, "public": public_dir})
    _write_docx(
        article_markdown,
        rich_docx_path,
        placements=placements,
        image_dirs={"generated": generated_dir, "public": public_dir},
    )

    result = {
        "contract_name": "publish_pack_result",
        "contract_version": "v1",
        "title": title,
        "platform": platform,
        "platform_name": _platform_name(platform),
        "output_dir": str(output_dir),
        "selected_image_plan": {
            "id": image_plan.item_id,
            "name": image_plan.title,
        }
        if image_plan
        else None,
        "image_brief": image_brief or None,
        "artifact_refs": [
            str(tasks_path),
            str(placements_path),
            str(sources_path),
            str(public_leads_path),
            str(pure_docx_path),
            str(rich_docx_path),
        ],
        "image_task_count": len(tasks),
        "images_present": _existing_image_count(slots, {"generated": generated_dir, "public": public_dir}),
        "recommended_next_actions": [
            "先补公开图片线索，至少保证 1-2 张不是生成图。",
            "检查 生成任务.json 的提示词是否足够纪实、真实，不要海报感。",
            "补齐或生成 图片/已生成 下的封面与配图后，可重跑 build-publish-pack 刷新图文 DOCX。",
            "如需真实生图，可再跑 render-packy-images。",
        ],
    }
    image_review_report = build_image_review_report(
        {
            "output_dir": str(output_dir),
            "platform": platform,
            "title": title,
            "topic": topic,
            "article_markdown": article_markdown,
            "context_pack": context_pack,
            "use_image_governance_review": payload.get("use_image_governance_review"),
        }
    )
    result["image_review_report"] = image_review_report
    review_actions = list(image_review_report.get("required_actions") or [])
    result["recommended_next_actions"] = list(
        dict.fromkeys([*review_actions, *result["recommended_next_actions"]])
    )
    return result


def _build_image_slots_from_brief(
    *,
    image_brief: dict[str, Any],
    title: str,
    topic: str,
    article_markdown: str,
    platform: str,
) -> list[ImageSlot]:
    paragraphs = [para for para in split_paragraphs(article_markdown) if not para.startswith("#")]
    fallback_anchors = _pick_anchor_paragraphs(paragraphs, max(1, len(image_brief.get("slots") or [])))
    slots: list[ImageSlot] = []
    for index, raw_slot in enumerate(image_brief.get("slots") or []):
        if not isinstance(raw_slot, dict):
            continue
        filename = str(raw_slot.get("filename") or ("封面" if index == 0 else f"配图-{index:02d}")).strip()
        role = str(raw_slot.get("role") or ("cover_opinion" if index == 0 else "supporting_visual")).strip()
        allowed_raw = [str(item).strip().lower() for item in raw_slot.get("allowed_source_types") or [] if str(item).strip()]
        if not allowed_raw:
            allowed_raw = [str(item).strip().lower() for item in raw_slot.get("source_priority") or [] if str(item).strip()]
        if not allowed_raw:
            allowed_raw = ["generated"] if index == 0 else ["public", "generated"]
        allowed = tuple(dict.fromkeys(item for item in allowed_raw if item in {"generated", "public"})) or ("generated",)
        primary_source_type = allowed[0]
        anchor = str(raw_slot.get("anchor") or (fallback_anchors[min(index, len(fallback_anchors) - 1)] if fallback_anchors else topic)).strip()
        caption = f"封面图：{compact_whitespace(title)}" if index == 0 else _caption_from_paragraph(anchor)
        prompt = _slot_prompt(
            topic=topic,
            snippet=anchor,
            style_hint="真实编辑感，避免海报感",
            source_type=primary_source_type,
        )
        aspect_ratio = "16:9" if index == 0 else ("4:3" if role != "product_screenshot" else "16:9")
        image_size = "2K" if index == 0 else "1K"
        slots.append(
            ImageSlot(
                filename=filename,
                caption=caption,
                prompt=prompt,
                aspect_ratio=aspect_ratio,
                image_size=image_size,
                source_type=primary_source_type,
                role=role,
                allowed_source_types=allowed,
                after_contains=anchor[:80] if index > 0 else None,
            )
        )
    return slots


def render_packy_images(payload: dict[str, Any]) -> dict[str, Any]:
    tasks_path = Path(str(payload.get("tasks_path") or "")).expanduser()
    output_dir = Path(str(payload.get("output_dir") or tasks_path.parent / "已生成")).expanduser()
    if not tasks_path.exists():
        raise ValueError("tasks_path 不存在")

    runtime = _resolve_image_runtime(payload)
    fallback_runtime = _resolve_image_fallback_runtime(payload, primary_runtime=runtime)
    if not runtime.api_key:
        return {
            "contract_name": "render_images_result",
            "contract_version": "v1",
            "status": "awaiting_api_key",
            "resolved_provider": runtime.provider,
            "resolved_model": runtime.model,
            "resolved_base_url": runtime.base_url,
            "artifact_refs": [],
            "recommended_next_actions": _missing_image_config_actions(runtime.provider),
        }

    output_dir.mkdir(parents=True, exist_ok=True)
    tasks = json.loads(tasks_path.read_text(encoding="utf-8"))
    retry_attempts = _resolve_image_retry_attempts(payload)
    request_model = _normalize_image_model_alias(runtime.model)
    fallback_request_model = _normalize_image_model_alias(fallback_runtime.model) if fallback_runtime else ""
    results: list[dict[str, Any]] = []
    for task in tasks:
        filename = str(task.get("filename") or "image").strip()
        prompt = str(task.get("prompt") or "").strip()
        if not prompt:
            results.append({"filename": filename, "status": "skipped", "error": "prompt 为空"})
            continue

        rendered = _call_image_api_with_retries(
            prompt=prompt,
            model=request_model,
            api_key=runtime.api_key,
            aspect_ratio=str(task.get("aspect_ratio") or "4:3"),
            image_size=str(task.get("image_size") or "1K"),
            base_url=runtime.base_url,
            retry_attempts=retry_attempts,
        )
        used_runtime = runtime
        used_request_model = request_model
        used_fallback = False
        if (
            rendered["status"] != "ok"
            and fallback_runtime is not None
            and _should_try_image_fallback(runtime, fallback_runtime)
        ):
            fallback_rendered = _call_image_api_with_retries(
                prompt=prompt,
                model=fallback_request_model,
                api_key=fallback_runtime.api_key,
                aspect_ratio=str(task.get("aspect_ratio") or "4:3"),
                image_size=str(task.get("image_size") or "1K"),
                base_url=fallback_runtime.base_url,
                retry_attempts=retry_attempts,
            )
            if fallback_rendered["status"] == "ok":
                rendered = fallback_rendered
                used_runtime = fallback_runtime
                used_request_model = fallback_request_model
                used_fallback = True
            else:
                rendered = {
                    "status": "error",
                    "error": f"primary={rendered.get('error') or 'unknown'} | fallback={fallback_rendered.get('error') or 'unknown'}",
                    "attempts_used": max(
                        int(rendered.get("attempts_used") or 0),
                        int(fallback_rendered.get("attempts_used") or 0),
                    ),
                }
        if rendered["status"] != "ok":
            results.append(
                {
                    "filename": filename,
                    "provider": used_runtime.provider,
                    "model": used_request_model,
                    "base_url": used_runtime.base_url,
                    "used_fallback": used_fallback,
                    **rendered,
                }
            )
            continue

        normalized_bytes, ext, width, height = _normalize_workflow_image_bytes(
            rendered["bytes"],
            extension=rendered["extension"],
            aspect_ratio=str(task.get("aspect_ratio") or "4:3"),
            image_size=str(task.get("image_size") or "1K"),
        )
        out_path = output_dir / f"{Path(filename).stem}.{ext}"
        out_path.write_bytes(normalized_bytes)
        _cleanup_slot_variants(output_dir, filename, keep_path=out_path)
        results.append(
            {
                "filename": filename,
                "status": "ok",
                "path": str(out_path),
                "mime_type": rendered["mime_type"],
                "width": width,
                "height": height,
                "provider": used_runtime.provider,
                "model": used_request_model,
                "base_url": used_runtime.base_url,
                "used_fallback": used_fallback,
            }
        )

    success_count = sum(1 for item in results if item["status"] == "ok")
    index_path = tasks_path.parent / GENERATED_INDEX_FILENAME
    _write_generated_index(
        index_path=index_path,
        tasks=tasks,
        results=results,
    )
    recommended_actions = _render_image_retry_actions(results, runtime, retry_attempts=retry_attempts)
    return {
        "contract_name": "render_images_result",
        "contract_version": "v1",
        "status": "completed" if success_count == len(results) else "partial",
        "resolved_provider": runtime.provider,
        "resolved_model": request_model,
        "configured_model": runtime.model,
        "resolved_base_url": runtime.base_url,
        "resolved_retry_attempts": retry_attempts,
        "fallback_provider": fallback_runtime.provider if fallback_runtime else "",
        "fallback_model": fallback_request_model if fallback_runtime else "",
        "fallback_base_url": fallback_runtime.base_url if fallback_runtime else "",
        "fallback_used": any(bool(item.get("used_fallback")) for item in results),
        "artifact_refs": [item["path"] for item in results if item["status"] == "ok"] + [str(index_path)],
        "results": results,
        "recommended_next_actions": recommended_actions,
    }


def _resolve_output_dir(payload: dict[str, Any], config: AppConfig, *, platform: str, title: str) -> Path:
    raw = str(payload.get("output_dir") or "").strip()
    if raw:
        path = Path(raw).expanduser()
        path.mkdir(parents=True, exist_ok=True)
        return path
    slug = slug_for_filename(title, "publish-pack")
    path = config.data_dir / "publish-packs" / slug / _platform_name(platform)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _select_image_plan(config: AppConfig, payload: dict[str, Any], context_pack: dict[str, Any], *, article_markdown: str):
    query_terms = extract_terms(
        str(payload.get("title") or ""),
        str(payload.get("topic") or context_pack.get("topic") or ""),
        str(payload.get("platform") or context_pack.get("platform") or ""),
        *[str(item) for item in context_pack.get("must_cover_points") or []],
    )
    plans = rank_items(load_markdown_items(config.image_plans_dir, "id", "name"), query_terms)
    if not plans:
        return None
    platform = str(payload.get("platform") or context_pack.get("platform") or "").strip().lower()
    if platform == "wechat" and _looks_like_opinion_article(str(payload.get("title") or ""), article_markdown):
        for plan in plans:
            if "观点型" in plan.title:
                return plan
    return plans[0]


def _build_image_slots(
    *,
    title: str,
    topic: str,
    article_markdown: str,
    platform: str,
    image_count: int,
    image_plan_name: str,
    image_plan_body: str,
) -> list[ImageSlot]:
    paragraphs = [para for para in split_paragraphs(article_markdown) if not para.startswith("#")]
    anchors = _pick_anchor_paragraphs(paragraphs, image_count)
    style_hint = _style_hint(image_plan_body)
    source_types = _slot_source_types(image_count, image_plan_name=image_plan_name, image_plan_body=image_plan_body)
    slots: list[ImageSlot] = [
        ImageSlot(
            filename="封面",
            caption=f"封面图：{compact_whitespace(title)}",
            prompt=(
                f"{topic}，中国语境，真实纪实摄影，新闻图片或杂志专题封面质感，适合作为{_platform_name(platform)}文章封面，"
                f"{style_hint}，自然光，真实人物或真实场景，不要 AI 海报感，不要赛博朋克，不要机器人，不出现品牌 logo"
            ),
            aspect_ratio="16:9",
            image_size="2K",
            source_type="generated",
            role="cover_opinion",
            allowed_source_types=("generated",),
        )
    ]
    for idx, paragraph in enumerate(anchors, 1):
        snippet = compact_whitespace(paragraph)
        source_type = source_types[idx - 1]
        slots.append(
            ImageSlot(
                filename=f"配图-{idx:02d}",
                caption=_caption_from_paragraph(snippet),
                prompt=_slot_prompt(
                    topic=topic,
                    snippet=snippet,
                    style_hint=style_hint,
                    source_type=source_type,
                ),
                aspect_ratio="4:3",
                image_size="1K",
                source_type=source_type,
                role="supporting_visual",
                allowed_source_types=(source_type,),
                after_contains=snippet[:80],
            )
        )
    return slots


def _pick_anchor_paragraphs(paragraphs: list[str], image_count: int) -> list[str]:
    if not paragraphs:
        return ["正文重点段落"] * image_count
    candidates = [para for para in paragraphs if len(compact_whitespace(para)) >= 24] or paragraphs
    total = len(candidates)
    slots: list[str] = []
    for idx in range(image_count):
        if total == 1:
            slots.append(candidates[0])
            continue
        anchor_index = round((idx + 1) * (total - 1) / (image_count + 1))
        slots.append(candidates[anchor_index])
    return slots


def _build_placements(slots: list[ImageSlot]) -> dict[str, Any]:
    cover = slots[0]
    inline = [
        {
            "path": _slot_relative_path(slot),
            "caption": slot.caption,
            "after_contains": slot.after_contains or "",
            "source_type": slot.source_type,
            "role": slot.role,
            "allowed_source_types": list(slot.allowed_source_types),
        }
        for slot in slots[1:]
    ]
    return {
        "cover": {
            "path": _slot_relative_path(cover),
            "caption": cover.caption,
            "source_type": cover.source_type,
            "role": cover.role,
            "allowed_source_types": list(cover.allowed_source_types),
        },
        "inline": inline,
    }


def _build_source_markdown(*, slots: list[ImageSlot], payload: dict[str, Any], image_model: str) -> str:
    public_refs = list(payload.get("public_image_refs") or [])
    lines = ["# 图片与数据来源", ""]
    for slot in slots:
        if slot.source_type == "generated" and "public" not in slot.allowed_source_types:
            lines.append(f"- {slot.filename}：使用 `packyapi` 图片接口生成")
        elif "public" in slot.allowed_source_types and "generated" in slot.allowed_source_types:
            lines.append(f"- {slot.filename}：优先使用公开真实图片，如无合适素材再考虑生成图补位")
        else:
            lines.append(f"- {slot.filename}：建议使用公开渠道真实图片，来源待补")
    if public_refs:
        lines.extend(["", "公开渠道补充图源："])
        for ref in public_refs:
            title = str(ref.get("title") or ref.get("label") or "未命名来源").strip()
            url = str(ref.get("url") or "").strip()
            note = str(ref.get("note") or "").strip()
            line = f"- {title}"
            if url:
                line += f"：{url}"
            if note:
                line += f"（{note}）"
            lines.append(line)

    lines.extend(
        [
            "",
            "说明：",
            "",
            "- 本文默认采用“生成图 + 公开渠道图”混搭，不建议全部使用 AI 生成图",
            f"- 本次生成使用的模型为 `{image_model}`",
            "- 如需公开渠道配图，请在 public_image_refs 中补充来源 URL、标题和使用说明",
        ]
    )
    return "\n".join(lines)


def _build_public_image_leads(slots: list[ImageSlot]) -> str:
    lines = ["# 公开图片线索", ""]
    public_slots = [slot for slot in slots if "public" in slot.allowed_source_types]
    if not public_slots:
        lines.append("- 当前批次没有公共图源槽位。")
        return "\n".join(lines)

    lines.append("建议优先使用公开新闻图片、会议现场图、公司发布会现场图、产品截图等真实感强的素材。")
    lines.append("")
    for slot in public_slots:
        lines.extend(
            [
                f"## {slot.filename}",
                f"- 用途：{slot.caption}",
                f"- 图位职责：{slot.role}",
                f"- 公开渠道搜索词：{slot.prompt}",
                f"- 建议来源：官方博客、官方新闻稿、Wikimedia Commons、公司 newsroom、产品发布页面",
                "- 待补字段：图片 URL、来源标题、许可/使用说明、下载后本地路径",
                "",
            ]
        )
    return "\n".join(lines).rstrip()


def _write_docx(article_markdown: str, path: Path, *, placements: dict[str, Any] | None, image_dirs: dict[str, Path]) -> None:
    title = _extract_title(article_markdown)
    paragraphs = split_paragraphs(article_markdown)
    body_paragraphs = [para for para in paragraphs if para != f"# {title}"]

    doc = Document()
    _configure_doc_styles(doc)
    _add_title_block(doc, title)
    if placements:
        _add_image_block(
            doc,
            _resolve_slot_image_path(str(placements["cover"]["path"]), image_dirs),
            cover=True,
        )

    for paragraph in body_paragraphs:
        _render_markdown_block(doc, paragraph)
        if not placements:
            continue
        for slot in placements.get("inline", []):
            anchor = str(slot.get("after_contains") or "").strip()
            if anchor and anchor in compact_whitespace(paragraph):
                _add_image_block(
                    doc,
                    _resolve_slot_image_path(str(slot.get("path") or ""), image_dirs),
                )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def _add_image_block(doc: Document, image_path: Path, *, cover: bool = False) -> None:
    if image_path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(10)
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(6.2 if cover else 5.6))


def _configure_doc_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.first_line_indent = Inches(0.28)

    title_style = doc.styles["Title"]
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.paragraph_format.space_after = Pt(14)

    heading1 = doc.styles["Heading 1"]
    heading1.font.size = Pt(15)
    heading1.font.bold = True
    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(8)

    heading2 = doc.styles["Heading 2"]
    heading2.font.size = Pt(13)
    heading2.font.bold = True
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(6)


def _add_title_block(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.add_run(title)


def _render_markdown_block(doc: Document, block: str) -> None:
    stripped = block.strip()
    if not stripped:
        return
    if stripped.startswith("## "):
        _add_heading_block(doc, stripped[3:].strip(), level=1)
        return
    if stripped.startswith("### "):
        _add_heading_block(doc, stripped[4:].strip(), level=2)
        return

    lines = [line.rstrip() for line in stripped.splitlines() if line.strip()]
    if lines and all(_is_bullet_line(line) for line in lines):
        for line in lines:
            _add_list_paragraph(doc, _strip_list_marker(line), numbered=False)
        return
    if lines and all(_is_numbered_line(line) for line in lines):
        for line in lines:
            _add_list_paragraph(doc, _strip_numbered_marker(line), numbered=True)
        return
    if lines and all(line.lstrip().startswith(">") for line in lines):
        _add_quote_block(doc, "\n".join(line.lstrip()[1:].strip() for line in lines))
        return

    _add_body_paragraph(doc, stripped)


def _add_heading_block(doc: Document, text: str, *, level: int) -> None:
    style_name = "Heading 1" if level <= 1 else "Heading 2"
    paragraph = doc.add_paragraph(style=style_name)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _append_inline_runs(paragraph, text)


def _add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    _append_inline_runs(paragraph, text)


def _add_list_paragraph(doc: Document, text: str, *, numbered: bool) -> None:
    paragraph = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    _append_inline_runs(paragraph, text)


def _add_quote_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.italic = True


def _append_inline_runs(paragraph, text: str) -> None:
    for token_type, token_text in _iter_inline_tokens(text):
        run = paragraph.add_run(token_text)
        if token_type == "bold":
            run.bold = True
        elif token_type == "code":
            run.font.name = "Menlo"
            run.font.size = Pt(10.5)


def _iter_inline_tokens(text: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    last_end = 0
    for match in pattern.finditer(text):
        if match.start() > last_end:
            tokens.append(("text", text[last_end:match.start()]))
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            tokens.append(("bold", token[2:-2]))
        elif token.startswith("`") and token.endswith("`"):
            tokens.append(("code", token[1:-1]))
        else:
            tokens.append(("text", token))
        last_end = match.end()
    if last_end < len(text):
        tokens.append(("text", text[last_end:]))
    return tokens or [("text", text)]


def _is_bullet_line(line: str) -> bool:
    stripped = line.lstrip()
    return stripped.startswith("- ") or stripped.startswith("* ")


def _is_numbered_line(line: str) -> bool:
    return bool(re.match(r"^\s*\d+\.\s+", line))


def _strip_list_marker(line: str) -> str:
    stripped = line.lstrip()
    return stripped[2:].strip() if len(stripped) >= 2 else stripped


def _strip_numbered_marker(line: str) -> str:
    return re.sub(r"^\s*\d+\.\s+", "", line).strip()


def _existing_image_count(slots: list[ImageSlot], image_dirs: dict[str, Path]) -> int:
    count = 0
    for slot in slots:
        for ext in ("png", "jpg", "jpeg", "webp"):
            roots = []
            for source_type in slot.allowed_source_types:
                roots.append(image_dirs["generated"] if source_type == "generated" else image_dirs["public"])
            for root in roots:
                if (root / f"{slot.filename}.{ext}").exists():
                    count += 1
                    break
            else:
                continue
            break
    return count


def _extract_title(article_markdown: str) -> str:
    for line in article_markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("# "):
            return stripped[2:].strip()
    for paragraph in split_paragraphs(article_markdown):
        if paragraph.strip():
            return compact_whitespace(paragraph)[:60]
    return "未命名文章"


def _platform_name(platform: str) -> str:
    return platform_display_name(platform)


def _platform_file_stem(platform: str) -> str:
    return platform_file_stem(platform)


def _style_hint(image_plan_body: str) -> str:
    if not image_plan_body.strip():
        return "媒体感，层次清楚"
    for line in image_plan_body.splitlines():
        stripped = line.strip().lstrip("- ").strip()
        if stripped.startswith("风格："):
            return stripped.removeprefix("风格：").strip()
    return "媒体感，层次清楚"


def _caption_from_paragraph(paragraph: str) -> str:
    snippet = paragraph[:36].rstrip("，。；：: ")
    return f"{snippet}。"


def _slot_source_types(image_count: int, *, image_plan_name: str, image_plan_body: str) -> list[str]:
    combined = f"{image_plan_name}\n{image_plan_body}"
    if any(token in combined for token in ["纯信息图", "纯结构图", "全生成", "全信息图"]):
        return ["generated"] * image_count
    if image_count <= 3:
        return ["public", "generated", "public"]
    sequence = ["public", "generated", "public", "generated", "public"]
    return sequence[:image_count]


def _looks_like_opinion_article(title: str, article_markdown: str) -> bool:
    combined = f"{title}\n{article_markdown}"
    opinion_markers = ["为什么", "不是", "而是", "我越来越觉得", "我越来越认同", "真正", "本质上", "更像是"]
    procedural_markers = ["步骤", "方法", "怎么做", "实操", "教程", "清单", "复盘", "第一", "第二", "第三"]
    opinion_score = sum(1 for marker in opinion_markers if marker in combined)
    procedural_score = sum(1 for marker in procedural_markers if marker in combined)
    return opinion_score >= procedural_score


def _slot_prompt(*, topic: str, snippet: str, style_hint: str, source_type: str) -> str:
    if source_type == "public":
        return _public_search_query(topic=topic, snippet=snippet)
    return (
        f"{topic}，围绕这段文字生成配图：{snippet[:120]}。"
        f"中国语境，真实纪实摄影，杂志专题图片质感，适合文章正文插图，{style_hint}，自然光，不要海报感，不要夸张 UI。"
    )


def _public_search_query(*, topic: str, snippet: str) -> str:
    combined = f"{topic} {snippet}"
    normalized = combined.lower()

    if any(token in normalized for token in ["老板", "企业", "预算", "管理", "协同", "流程", "组织"]):
        return "business meeting executives office teamwork"
    if any(token in normalized for token in ["agent", "ai", "人工智能", "模型", "prompt", "智能体"]):
        return "AI 人工智能 会议 现场 商务"
    if any(token in normalized for token in ["团队", "协作", "办公室", "会议"]):
        return "technology team office meeting"

    short_topic = _compact_public_query_fragment(topic, limit=12)
    if short_topic:
        return f"{short_topic} 商务 团队 会议 真实照片"[:72]
    return "business office teamwork meeting"


def _compact_public_query_fragment(text: str, *, limit: int) -> str:
    normalized = re.sub(r"[*_`#>\[\]\(\)]+", "", text or "")
    normalized = re.sub(r"\s+", "", normalized)
    for sep in ["：", "，", "。", "；", "、", ":", ","]:
        normalized = normalized.split(sep, 1)[0]
    return normalized[:limit].strip()


def _slot_relative_path(slot: ImageSlot) -> str:
    root = "已生成" if slot.source_type == "generated" else "公开来源"
    return f"{root}/{slot.filename}.png"


def _resolve_slot_image_path(raw_path: str, image_dirs: dict[str, Path]) -> Path:
    path = Path(raw_path)
    root_name = path.parts[0] if path.parts else "已生成"
    root = image_dirs["public"] if root_name == "公开来源" else image_dirs["generated"]
    direct = root / path.name
    if direct.exists():
        return direct
    stem = path.stem
    for ext in ("png", "jpg", "jpeg", "webp"):
        candidate = root / f"{stem}.{ext}"
        if candidate.exists():
            return candidate
    alternate_root = image_dirs["public"] if root == image_dirs["generated"] else image_dirs["generated"]
    for ext in ("png", "jpg", "jpeg", "webp"):
        candidate = alternate_root / f"{stem}.{ext}"
        if candidate.exists():
            return candidate
    return direct


def _normalize_slot_assets(slots: list[ImageSlot], image_dirs: dict[str, Path]) -> None:
    for slot in slots:
        image_path = _resolve_slot_image_path(_slot_relative_path(slot), image_dirs)
        if not image_path.exists():
            continue
        normalized_path = _normalize_workflow_image_file(
            image_path,
            aspect_ratio=slot.aspect_ratio,
            image_size=slot.image_size,
        )
        _cleanup_slot_variants(normalized_path.parent, normalized_path.stem, keep_path=normalized_path)


def _normalize_workflow_image_file(path: Path, *, aspect_ratio: str, image_size: str) -> Path:
    raw = path.read_bytes()
    normalized_bytes, normalized_ext, _, _ = _normalize_workflow_image_bytes(
        raw,
        extension=path.suffix,
        aspect_ratio=aspect_ratio,
        image_size=image_size,
    )
    normalized_path = path.with_suffix("." + normalized_ext)
    normalized_path.write_bytes(normalized_bytes)
    if normalized_path != path and path.exists():
        path.unlink()
    return normalized_path


def _normalize_workflow_image_bytes(
    raw: bytes,
    *,
    extension: str,
    aspect_ratio: str,
    image_size: str,
) -> tuple[bytes, str, int, int]:
    target_width, target_height = _target_image_dimensions(aspect_ratio, image_size)
    normalized_ext = _normalized_image_extension(extension)
    try:
        with Image.open(io.BytesIO(raw)) as source:
            image = ImageOps.exif_transpose(source)
            fitted = ImageOps.fit(
                image,
                (target_width, target_height),
                method=Image.Resampling.LANCZOS,
                centering=(0.5, 0.5),
            )
            output = io.BytesIO()
            if normalized_ext == "jpg":
                fitted.convert("RGB").save(output, format="JPEG", quality=90, optimize=True)
            else:
                if fitted.mode not in {"RGB", "RGBA"}:
                    fitted = fitted.convert("RGBA" if "A" in fitted.getbands() else "RGB")
                fitted.save(output, format="PNG", optimize=True)
            return output.getvalue(), normalized_ext, target_width, target_height
    except (UnidentifiedImageError, OSError, ValueError):
        return raw, normalized_ext, target_width, target_height


def _target_image_dimensions(aspect_ratio: str, image_size: str) -> tuple[int, int]:
    normalized = str(aspect_ratio or "").strip()
    if normalized in STANDARD_IMAGE_TARGETS:
        return STANDARD_IMAGE_TARGETS[normalized]
    match = re.match(r"^\s*(\d+)\s*:\s*(\d+)\s*$", normalized)
    if not match:
        return STANDARD_IMAGE_TARGETS["4:3"]
    width_ratio = max(int(match.group(1)), 1)
    height_ratio = max(int(match.group(2)), 1)
    long_edge = DEFAULT_STANDARD_IMAGE_LONG_EDGE
    if width_ratio >= height_ratio:
        width = long_edge
        height = max(1, round(long_edge * height_ratio / width_ratio))
    else:
        height = long_edge
        width = max(1, round(long_edge * width_ratio / height_ratio))
    return width, height


def _normalized_image_extension(extension: str) -> str:
    normalized = extension.strip().lower().lstrip(".").replace("jpeg", "jpg")
    if normalized == "jpg":
        return "jpg"
    if normalized == "png":
        return "png"
    return "png"


def _cleanup_slot_variants(directory: Path, stem: str, *, keep_path: Path) -> None:
    for ext in ("png", "jpg", "jpeg", "webp"):
        candidate = directory / f"{stem}.{ext}"
        if candidate == keep_path:
            continue
        if candidate.exists():
            candidate.unlink()


def _write_generated_index(
    *,
    index_path: Path,
    tasks: list[dict[str, Any]],
    results: list[dict[str, Any]],
) -> None:
    task_lookup = {
        str(task.get("filename") or "").strip(): dict(task)
        for task in tasks
        if str(task.get("filename") or "").strip()
    }
    index = read_image_index(index_path)
    for item in results:
        if item.get("status") != "ok":
            continue
        filename = str(item.get("filename") or "").strip()
        task = task_lookup.get(filename, {})
        index[filename] = {
            "path": str(item.get("path") or ""),
            "source_type": "generated",
            "provider": str(item.get("provider") or ""),
            "model": str(item.get("model") or ""),
            "prompt": str(task.get("prompt") or ""),
            "aspect_ratio": str(task.get("aspect_ratio") or ""),
            "image_size": str(task.get("image_size") or ""),
            "mime_type": str(item.get("mime_type") or ""),
            "width": int(item.get("width") or 0),
            "height": int(item.get("height") or 0),
            "publishable": True,
            "publishability_reason": "generated_asset",
        }
    write_image_index(index_path, index)

def _call_image_api(
    *,
    prompt: str,
    model: str,
    api_key: str,
    aspect_ratio: str,
    image_size: str,
    base_url: str | None = None,
) -> dict[str, Any]:
    raw_base_url = base_url or _resolve_image_base_url()
    if _use_gemini_native_image_api(raw_base_url, model):
        return _call_gemini_native_image_api(
            prompt=prompt,
            model=model,
            api_key=api_key,
            aspect_ratio=aspect_ratio,
            image_size=image_size,
            base_url=raw_base_url,
        )
    return _call_openai_image_api(
        prompt=prompt,
        model=model,
        api_key=api_key,
        aspect_ratio=aspect_ratio,
        image_size=image_size,
        base_url=raw_base_url,
    )


def _call_openai_image_api(
    *,
    prompt: str,
    model: str,
    api_key: str,
    aspect_ratio: str,
    image_size: str,
    base_url: str,
) -> dict[str, Any]:
    base_url = normalize_base_url(base_url)
    payload = {
        "model": model,
        "prompt": prompt,
        "size": image_size,
        "aspect_ratio": aspect_ratio,
        "response_format": "b64_json",
    }
    command = [
        "curl",
        "-sS",
        "--max-time",
        "180",
        f"{base_url}/images/generations",
        "-H",
        "Content-Type: application/json",
        "-H",
        f"Authorization: Bearer {api_key}",
        "--data",
        json.dumps(payload, ensure_ascii=False),
    ]
    try:
        completed = subprocess.run(command, check=False, capture_output=True, text=True, timeout=180)
    except Exception as exc:
        return {"status": "error", "error": repr(exc)}
    if completed.returncode != 0:
        detail = completed.stderr.strip() or completed.stdout.strip() or f"exit={completed.returncode}"
        return {"status": "error", "error": detail[:400]}

    try:
        body = json.loads(completed.stdout)
    except Exception as exc:
        return {"status": "error", "error": f"invalid_json: {exc!r}"}
    if "error" in body:
        return {"status": "error", "error": json.dumps(body["error"], ensure_ascii=False)}
    data = body.get("data") or []
    if not data:
        return {"status": "error", "error": "empty_data"}
    item = data[0]
    b64 = str(item.get("b64_json") or "").strip()
    url = str(item.get("url") or "").strip()
    if b64:
        raw = base64.b64decode(b64)
        return {
            "status": "ok",
            "bytes": raw,
            "extension": _guess_image_extension(raw),
            "mime_type": "image/png",
        }
    if url:
        fetch = subprocess.run(
            ["curl", "-sS", "--max-time", "180", url],
            check=False,
            capture_output=True,
            timeout=180,
        )
        if fetch.returncode != 0:
            detail = (fetch.stderr or fetch.stdout).decode("utf-8", errors="ignore").strip()
            return {"status": "error", "error": detail[:400]}
        raw = fetch.stdout
        return {
            "status": "ok",
            "bytes": raw,
            "extension": _guess_image_extension(raw),
            "mime_type": "image/png",
        }
    return {"status": "error", "error": "missing_b64_or_url"}


def _call_gemini_native_image_api(
    *,
    prompt: str,
    model: str,
    api_key: str,
    aspect_ratio: str,
    image_size: str,
    base_url: str,
) -> dict[str, Any]:
    normalized_base_url = normalize_base_url(base_url)
    if normalized_base_url.endswith("/v1"):
        normalized_base_url = normalized_base_url[: -len("/v1")]
    endpoint = f"{normalized_base_url.rstrip('/')}/v1beta/models/{model}:generateContent"
    payload = {
        "contents": [
            {
                "parts": [
                    {"text": prompt},
                ]
            }
        ],
        "generationConfig": {
            "responseModalities": ["IMAGE"],
            "imageConfig": {
                "aspectRatio": aspect_ratio,
                "image_size": image_size,
            },
        },
    }
    command = [
        "curl",
        "-sS",
        "--max-time",
        "600",
        endpoint,
        "-H",
        "Content-Type: application/json",
        "-H",
        f"Authorization: Bearer {api_key}",
        "--data",
        json.dumps(payload, ensure_ascii=False),
    ]
    try:
        completed = subprocess.run(command, check=False, capture_output=True, text=True, timeout=600)
    except Exception as exc:
        return {"status": "error", "error": repr(exc)}
    if completed.returncode != 0:
        detail = completed.stderr.strip() or completed.stdout.strip() or f"exit={completed.returncode}"
        return {"status": "error", "error": detail[:400]}

    try:
        body = json.loads(completed.stdout)
    except Exception as exc:
        return {"status": "error", "error": f"invalid_json: {exc!r}"}
    if "error" in body:
        return {"status": "error", "error": json.dumps(body["error"], ensure_ascii=False)}

    candidate_summaries: list[str] = []
    for candidate in body.get("candidates") or []:
        finish_reason = str(candidate.get("finishReason") or candidate.get("finish_reason") or "").strip() or "unknown"
        content = candidate.get("content") or {}
        parts = content.get("parts") or []
        if not parts:
            candidate_summaries.append(f"finish_reason={finish_reason},parts=0")
            continue
        for part in parts:
            inline = part.get("inlineData") or {}
            b64 = str(inline.get("data") or "").strip()
            if b64:
                raw = base64.b64decode(b64)
                mime_type = str(inline.get("mimeType") or "image/png")
                return {
                    "status": "ok",
                    "bytes": raw,
                    "extension": _mime_to_extension(mime_type, raw),
                    "mime_type": mime_type,
                }
            # Fallback: some relays return base64 as a data URI inside the text field
            text_val = str(part.get("text") or "").strip()
            data_uri_match = re.search(r"data:(image/[a-zA-Z+]+);base64,([A-Za-z0-9+/=\s]+)", text_val)
            if data_uri_match:
                mime_type = data_uri_match.group(1)
                raw = base64.b64decode(data_uri_match.group(2))
                return {
                    "status": "ok",
                    "bytes": raw,
                    "extension": _mime_to_extension(mime_type, raw),
                    "mime_type": mime_type,
                }
        candidate_summaries.append(f"finish_reason={finish_reason},parts={len(parts)}")
    if candidate_summaries:
        return {"status": "error", "error": f"empty_candidate_parts: {'; '.join(candidate_summaries)}"}
    return {"status": "error", "error": "missing_candidates"}


def _guess_image_extension(raw: bytes) -> str:
    if raw.startswith(b"\x89PNG"):
        return "png"
    if raw.startswith(b"\xff\xd8"):
        return "jpg"
    if raw.startswith(b"RIFF") and b"WEBP" in raw[:16]:
        return "webp"
    return "png"


def _mime_to_extension(mime_type: str, raw: bytes) -> str:
    normalized = mime_type.strip().lower()
    if normalized == "image/jpeg":
        return "jpg"
    if normalized == "image/webp":
        return "webp"
    if normalized == "image/png":
        return "png"
    return _guess_image_extension(raw)


def _use_gemini_native_image_api(base_url: str, model: str) -> bool:
    normalized_base = base_url.lower()
    normalized_model = _normalize_image_model_alias(model)
    return "ikuncode.cc" in normalized_base and normalized_model in {
        "gemini-3.1-flash-image-preview",
        "gemini-3-pro-image-preview",
        "gemini-2.5-flash-image",
    }


def _resolve_image_api_key(payload: dict[str, Any]) -> str:
    return _resolve_image_runtime(payload).api_key


def _resolve_image_model(payload: dict[str, Any]) -> str:
    return _resolve_image_runtime(payload).model


def _resolve_image_base_url(provider: str | None = None) -> str:
    return _resolve_image_runtime({"image_provider": provider} if provider else {}).base_url


def _resolve_image_provider(payload: dict[str, Any]) -> str:
    return _resolve_image_runtime(payload).provider


def _resolve_image_runtime(payload: dict[str, Any]) -> ImageRuntime:
    explicit_provider = str(payload.get("image_provider") or "").strip().lower()
    explicit_api_key = str(payload.get("image_api_key") or "").strip()
    explicit_model = str(payload.get("image_model") or "").strip()
    explicit_base_url = str(payload.get("image_base_url") or "").strip()

    if explicit_provider == "packyapi" or os.environ.get("PACKYAPI_IMAGE_API_KEY"):
        return _resolve_packy_image_runtime(
            explicit_api_key=explicit_api_key,
            explicit_model=explicit_model,
            explicit_base_url=explicit_base_url,
        )
    if explicit_provider == "ikun" or explicit_provider == "image-relay":
        return _resolve_ikun_image_runtime(
            explicit_api_key=explicit_api_key,
            explicit_model=explicit_model,
            explicit_base_url=explicit_base_url,
        )
    if explicit_api_key or explicit_model or explicit_base_url:
        provider = _infer_image_provider(base_url=explicit_base_url, model=explicit_model)
        if provider == "packyapi":
            return _resolve_packy_image_runtime(
                explicit_api_key=explicit_api_key,
                explicit_model=explicit_model,
                explicit_base_url=explicit_base_url,
            )
        return _resolve_ikun_image_runtime(
            explicit_api_key=explicit_api_key,
            explicit_model=explicit_model,
            explicit_base_url=explicit_base_url,
        )
    if os.environ.get("IKUN_IMAGE_API_KEY") or os.environ.get("WRITING_BRAIN_IMAGE_API_KEY") or os.environ.get("IKUN_API_KEY"):
        return _resolve_ikun_image_runtime(
            explicit_api_key="",
            explicit_model="",
            explicit_base_url="",
        )
    return _resolve_ikun_image_runtime(
        explicit_api_key="",
        explicit_model="",
        explicit_base_url="",
    )


def _resolve_image_fallback_runtime(payload: dict[str, Any], *, primary_runtime: ImageRuntime) -> ImageRuntime | None:
    explicit_provider = str(payload.get("image_fallback_provider") or os.environ.get("IMAGE_FALLBACK_PROVIDER") or "").strip().lower()
    explicit_api_key = str(payload.get("image_fallback_api_key") or os.environ.get("IMAGE_FALLBACK_API_KEY") or "").strip()
    explicit_model = str(payload.get("image_fallback_model") or os.environ.get("IMAGE_FALLBACK_MODEL") or primary_runtime.model).strip()
    explicit_base_url = str(payload.get("image_fallback_base_url") or os.environ.get("IMAGE_FALLBACK_BASE_URL") or "").strip()

    if not explicit_provider and not explicit_api_key and not explicit_base_url:
        return None

    provider = explicit_provider or _infer_image_provider(base_url=explicit_base_url, model=explicit_model)
    if provider == "packyapi":
        runtime = _resolve_packy_image_runtime(
            explicit_api_key=explicit_api_key,
            explicit_model=explicit_model,
            explicit_base_url=explicit_base_url or DEFAULT_PACKY_IMAGE_BASE_URL,
        )
    else:
        runtime = _resolve_ikun_image_runtime(
            explicit_api_key=explicit_api_key,
            explicit_model=explicit_model,
            explicit_base_url=explicit_base_url or DEFAULT_IKUN_IMAGE_BASE_URL,
        )

    if not runtime.api_key:
        return None
    return runtime


def _resolve_packy_image_runtime(*, explicit_api_key: str, explicit_model: str, explicit_base_url: str) -> ImageRuntime:
    api_key, key_source = _first_present(
        ("payload.image_api_key", explicit_api_key),
        ("PACKYAPI_IMAGE_API_KEY", os.environ.get("PACKYAPI_IMAGE_API_KEY") or ""),
    )
    model, model_source = _first_present(
        ("payload.image_model", explicit_model),
        ("PACKYAPI_IMAGE_MODEL", os.environ.get("PACKYAPI_IMAGE_MODEL") or ""),
        ("default", DEFAULT_PACKY_IMAGE_MODEL),
    )
    base_url, base_url_source = _first_present(
        ("payload.image_base_url", explicit_base_url),
        ("PACKYAPI_IMAGE_BASE_URL", os.environ.get("PACKYAPI_IMAGE_BASE_URL") or ""),
        ("default", DEFAULT_PACKY_IMAGE_BASE_URL),
    )
    return ImageRuntime(
        provider="packyapi",
        api_key=api_key,
        model=model,
        base_url=_normalize_image_base_url(base_url, provider="packyapi"),
        key_source=key_source,
        model_source=model_source,
        base_url_source=base_url_source,
    )


def _resolve_ikun_image_runtime(*, explicit_api_key: str, explicit_model: str, explicit_base_url: str) -> ImageRuntime:
    api_key, key_source = _first_present(
        ("payload.image_api_key", explicit_api_key),
        ("IKUN_IMAGE_API_KEY", os.environ.get("IKUN_IMAGE_API_KEY") or ""),
        ("WRITING_BRAIN_IMAGE_API_KEY", os.environ.get("WRITING_BRAIN_IMAGE_API_KEY") or ""),
        ("IKUN_API_KEY", os.environ.get("IKUN_API_KEY") or ""),
    )
    model, model_source = _first_present(
        ("payload.image_model", explicit_model),
        ("IKUN_IMAGE_MODEL", os.environ.get("IKUN_IMAGE_MODEL") or ""),
        ("WRITING_BRAIN_IMAGE_MODEL", os.environ.get("WRITING_BRAIN_IMAGE_MODEL") or ""),
        ("default", DEFAULT_IMAGE_MODEL),
    )
    base_url, base_url_source = _first_present(
        ("payload.image_base_url", explicit_base_url),
        ("IKUN_IMAGE_BASE_URL", os.environ.get("IKUN_IMAGE_BASE_URL") or ""),
        ("WRITING_BRAIN_IMAGE_BASE_URL", os.environ.get("WRITING_BRAIN_IMAGE_BASE_URL") or ""),
        ("IKUN_BASE_URL", os.environ.get("IKUN_BASE_URL") or ""),
        ("default", DEFAULT_IKUN_IMAGE_BASE_URL),
    )
    return ImageRuntime(
        provider="ikun",
        api_key=api_key,
        model=model,
        base_url=_normalize_image_base_url(base_url, provider="ikun"),
        key_source=key_source,
        model_source=model_source,
        base_url_source=base_url_source,
    )


def _infer_image_provider(*, base_url: str, model: str) -> str:
    lowered_base = base_url.strip().lower()
    lowered_model = model.strip().lower()
    if "ikuncode.cc" in lowered_base:
        return "ikun"
    if "packyapi" in lowered_base or lowered_model.startswith("gpt-image-"):
        return "packyapi"
    return "ikun"


def _should_try_image_fallback(primary_runtime: ImageRuntime, fallback_runtime: ImageRuntime) -> bool:
    return (
        primary_runtime.provider != fallback_runtime.provider
        or primary_runtime.base_url != fallback_runtime.base_url
        or _normalize_image_model_alias(primary_runtime.model) != _normalize_image_model_alias(fallback_runtime.model)
        or primary_runtime.api_key != fallback_runtime.api_key
    )


def _normalize_image_base_url(base_url: str, *, provider: str) -> str:
    normalized = normalize_base_url(base_url)
    if provider == "packyapi" and normalized.endswith("/v1"):
        return normalized[: -len("/v1")]
    return normalized


def _first_present(*items: tuple[str, str]) -> tuple[str, str]:
    for source, value in items:
        cleaned = value.strip()
        if cleaned:
            return cleaned, source
    return "", ""


def _normalize_image_model_alias(model: str) -> str:
    normalized = model.strip().lower()
    return IMAGE_MODEL_ALIASES.get(normalized, normalized)


def _missing_image_config_actions(provider: str) -> list[str]:
    if provider == "packyapi":
        return [
            "配置 PACKYAPI_IMAGE_API_KEY 后重试。",
            "如需自定义路由，可同时配置 PACKYAPI_IMAGE_BASE_URL；默认将使用 https://www.packyapi.com。",
            "如未指定模型，默认将使用 gemini-3.1-flash-image-preview。",
        ]
    return [
        "配置 IKUN_IMAGE_API_KEY 后重试。",
        "如需兼容旧配置，可继续使用 WRITING_BRAIN_IMAGE_API_KEY 或 IKUN_API_KEY。",
        "如未指定模型，默认将使用 nano2（实际请求 gemini-3.1-flash-image-preview）。",
    ]


def _render_image_retry_actions(results: list[dict[str, Any]], runtime: ImageRuntime, *, retry_attempts: int) -> list[str]:
    success_count = sum(1 for item in results if item.get("status") == "ok")
    if success_count:
        return ["图片生成完成后，重跑 build-publish-pack 刷新图文 DOCX。"]

    combined_errors = " ".join(str(item.get("error") or "") for item in results).lower()
    if "无效的令牌" in combined_errors or "invalid" in combined_errors and "token" in combined_errors:
        if runtime.provider == "packyapi":
            return [
                "当前 PACKYAPI 图片令牌无效；请改配 PACKYAPI_IMAGE_API_KEY，并确认它来自官方文档要求的 gemini 分组。",
                "图片请求将使用 https://www.packyapi.com；不要直接复用聊天通道的 /v1 路由。",
            ]
        return [
            "当前 IKUN 图片令牌无效；请更新 IKUN_IMAGE_API_KEY 或 IKUN_API_KEY 后重试。",
        ]
    if "分组" in combined_errors or "not supported model for image generation" in combined_errors or "model_not_found" in combined_errors:
        if runtime.provider == "packyapi":
            return [
                "当前 PACKYAPI 图片令牌与模型分组不匹配；请使用 PACKYAPI_IMAGE_API_KEY 单独配置图片通道。",
                "按照官方教程，绘图令牌应来自 gemini 分组，API 地址使用 https://www.packyapi.com。",
            ]
        return [
            f"当前 IKUN 图片模型不可用；系统已按同一模型重试 {retry_attempts} 次，请检查 IKUN_IMAGE_MODEL 或当前账号分组是否支持该模型。",
        ]
    if "empty_candidate_parts" in combined_errors:
        return [
            "IKUN 图像接口返回了成功状态，但 candidates.content.parts 为空；这与官方文档示例不一致，通常是中转通道未真正产出图片。",
            "先确认当前令牌组确实支持 NanoBanana 图像模型；如果仍复现，需要把原始响应提交给 IKUN 侧排查。",
        ]
    return ["先处理失败的图片任务，再重跑 render-packy-images。"]


def _resolve_image_retry_attempts(payload: dict[str, Any]) -> int:
    raw = str(
        payload.get("image_retry_attempts")
        or os.environ.get("WRITING_BRAIN_IMAGE_RETRY_ATTEMPTS")
        or DEFAULT_IMAGE_RETRY_ATTEMPTS
    ).strip()
    try:
        value = int(raw)
    except Exception:
        value = DEFAULT_IMAGE_RETRY_ATTEMPTS
    return max(1, value)


def _call_image_api_with_retries(
    *,
    prompt: str,
    model: str,
    api_key: str,
    aspect_ratio: str,
    image_size: str,
    base_url: str,
    retry_attempts: int,
) -> dict[str, Any]:
    attempts = max(1, retry_attempts)
    last_result: dict[str, Any] = {"status": "error", "error": "unknown"}
    for attempt in range(1, attempts + 1):
        result = _call_image_api(
            prompt=prompt,
            model=model,
            api_key=api_key,
            aspect_ratio=aspect_ratio,
            image_size=image_size,
            base_url=base_url,
        )
        if result.get("status") == "ok":
            return result
        last_result = dict(result)
        last_result["attempts_used"] = attempt
    return last_result
